import json
import zipfile
from io import BytesIO
from pathlib import Path

import openpyxl
from pypdf import PdfReader
from reportlab.pdfgen import canvas

import manufacturing_package as mp
import vault_tools_cli as cli

YELLOW = "FFF59D"


def _pdf(path: Path, label: str, pages: int = 1) -> Path:
    c = canvas.Canvas(str(path), pagesize=(1224, 792))
    for p in range(pages):
        c.drawString(100, 700, f"{label} sheet {p + 1}")
        c.showPage()
    c.save()
    return path


def _partner_manifest(tmp_path: Path, **over) -> Path:
    stp = tmp_path / "001.stp"; stp.write_bytes(b"ISO-10303-21;")
    released = _pdf(tmp_path / "002.pdf", "PLATE")
    draft = _pdf(tmp_path / "003.pdf", "BRACKET", pages=2)
    m = {
        "kind": "partner", "assembly": "CD-001882", "description": "Big blue adapter", "revision": "A",
        "state": "Released", "partner": "Acme Machining", "mode": "CNC",
        "source": "Model state: Primary | Design view: Default | BOM view: Structured",
        "user": "zak", "created": "2026-10-13 14:05",
        "rows": [
            {"item": "1", "part_number": "CD-000123", "description": "Plate", "qty": 4, "material": "Aluminum 6061",
             "finish": "As Finished", "process": "CNC", "revision": "A", "state": "Released",
             "files": ["STEP/CD-000123 Rev A.stp", "PDF/CD-000123 Rev A.pdf"]},
            {"item": "2", "part_number": "CD-000124", "description": "Bracket", "qty": 1, "material": "Steel",
             "finish": "", "process": "guessed: CNC", "revision": "", "state": "Work in Progress",
             "files": ["PDF/CD-000124 Rev -.pdf"]},
        ],
        "files": [
            {"src": str(stp), "dest": "STEP/CD-000123 Rev A.stp", "stamp": False},
            {"src": str(released), "dest": "PDF/CD-000123 Rev A.pdf", "stamp": False},
            {"src": str(draft), "dest": "PDF/CD-000124 Rev -.pdf", "stamp": True},
        ],
        "guessed": ["CD-000124 (CNC)"], "missing_drawings": ["CD-000125"], "no_flat_pattern": [],
        "failed": [{"part": "CD-000126", "reason": "STEP export failed: boom"}],
    }
    m.update(over)
    p = tmp_path / "manifest.json"
    p.write_text(json.dumps(m), encoding="utf-8")
    return p


def test_partner_zip_layout_and_counts(tmp_path):
    out = tmp_path / "CD-001882 Acme Machining CNC Package.zip"
    res = mp.run_package(str(_partner_manifest(tmp_path)), str(out))
    assert res["ok"] is True, res
    names = zipfile.ZipFile(out).namelist()
    assert names[:2] == ["00 Cover.pdf", "BOM.xlsx"]
    assert set(names) == {"00 Cover.pdf", "BOM.xlsx", "STEP/CD-000123 Rev A.stp",
                          "PDF/CD-000123 Rev A.pdf", "PDF/CD-000124 Rev -.pdf"}
    assert not any(n.endswith("/") for n in names) and not any(n.startswith("DXF") for n in names)
    assert (res["step"], res["pdf"], res["dxf"], res["print"], res["stamped"]) == (1, 2, 0, 0, 1)
    assert res["files"] == 5 and res["rows"] == 2


def test_partner_bom_xlsx_title_source_state_and_yellow(tmp_path):
    out = tmp_path / "p.zip"
    mp.run_package(str(_partner_manifest(tmp_path)), str(out))
    ws = openpyxl.load_workbook(BytesIO(zipfile.ZipFile(out).read("BOM.xlsx")))["BOM"]
    assert ws["A1"].value == "CD-001882 Acme Machining CNC package"
    assert ws["A2"].value.startswith("Generated: ") and ws["A2"].value.endswith(" | Model state: Primary | Design view: Default | BOM view: Structured")
    assert [c.value for c in ws[3]] == ["Item", "Part Number", "Description", "Qty", "Material", "Material Finish",
                                        "Process", "REV", "State", "File(s) in package"]
    assert ws["G5"].value == "guessed: CNC" and ws["I5"].value == "Work in Progress"
    assert ws["J4"].value == "CD-000123 Rev A.stp, CD-000123 Rev A.pdf"
    assert (ws["A4"].fill.fgColor.rgb or "")[-6:] != YELLOW
    assert (ws["A5"].fill.fgColor.rgb or "")[-6:] == YELLOW
    assert ws["D4"].value == 4


def test_stamp_only_flagged_pdfs(tmp_path):
    out = tmp_path / "p.zip"
    mp.run_package(str(_partner_manifest(tmp_path)), str(out))
    z = zipfile.ZipFile(out)
    stamped = PdfReader(BytesIO(z.read("PDF/CD-000124 Rev -.pdf")))
    plain = PdfReader(BytesIO(z.read("PDF/CD-000123 Rev A.pdf")))
    assert all("NOT RELEASED" in (p.extract_text() or "") for p in stamped.pages) and len(stamped.pages) == 2
    assert all("NOT RELEASED" not in (p.extract_text() or "") for p in plain.pages)


def test_cover_lists(tmp_path):
    out = tmp_path / "p.zip"
    mp.run_package(str(_partner_manifest(tmp_path)), str(out))
    text = "".join(p.extract_text() or "" for p in PdfReader(BytesIO(zipfile.ZipFile(out).read("00 Cover.pdf"))).pages)
    for s in ("Partner package: CD-001882", "Partner: Acme Machining", "Parts: CNC (machined)",
              "Rev A (Released)", "Process not set (guessed)", "CD-000124 (CNC)", "Missing drawings", "CD-000125",
              "No flat pattern", "Could not export", "CD-000126: STEP export failed: boom", "2 parts: 1 STEP, 2 PDF, 0 DXF"):
        assert s in text, s


def test_missing_src_goes_to_failed_and_out_of_row(tmp_path):
    m = json.loads(_partner_manifest(tmp_path).read_text(encoding="utf-8"))
    m["files"][0]["src"] = str(tmp_path / "gone.stp")
    p = tmp_path / "m2.json"; p.write_text(json.dumps(m), encoding="utf-8")
    out = tmp_path / "p.zip"
    res = mp.run_package(str(p), str(out))
    assert res["ok"] is True and res["step"] == 0 and res["failed"] == 2
    ws = openpyxl.load_workbook(BytesIO(zipfile.ZipFile(out).read("BOM.xlsx")))["BOM"]
    assert ws["J4"].value == "CD-000123 Rev A.pdf"


def test_bad_dest_is_rejected_and_no_zip(tmp_path):
    for bad in ("../evil.stp", "C:/x.stp", "/abs.stp", "STEP/../../x.stp"):
        m = json.loads(_partner_manifest(tmp_path).read_text(encoding="utf-8"))
        m["files"][0]["dest"] = bad
        p = tmp_path / "bad.json"; p.write_text(json.dumps(m), encoding="utf-8")
        out = tmp_path / "bad.zip"
        res = mp.run_package(str(p), str(out))
        assert res["ok"] is False and "Bad package path" in res["error"] and not out.exists()


def test_print_package_layout_and_list(tmp_path):
    f = tmp_path / "a.3mf"; f.write_bytes(b"3mf")
    m = {"kind": "print", "assembly": "CD-001882", "mode": "3D Print",
         "rows": [{"item": "3", "part_number": "CD-000200", "description": "Clip", "qty": 4, "material": "PLA",
                   "process": "3D Print", "revision": "", "state": "", "files": ["PLA/CD-000200 Rev - x4.3mf"]}],
         "files": [{"src": str(f), "dest": "PLA/CD-000200 Rev - x4.3mf", "stamp": False}]}
    p = tmp_path / "m.json"; p.write_text(json.dumps(m), encoding="utf-8")
    out = tmp_path / "CD-001882 3D Print Package.zip"
    res = mp.run_package(str(p), str(out))
    assert res["ok"] is True and res["print"] == 1
    names = zipfile.ZipFile(out).namelist()
    assert names == ["Print List.xlsx", "PLA/CD-000200 Rev - x4.3mf"]
    ws = openpyxl.load_workbook(BytesIO(zipfile.ZipFile(out).read("Print List.xlsx")))["Print List"]
    assert ws["A1"].value == "CD-001882 3D print list"
    assert "|" not in ws["A2"].value
    assert [c.value for c in ws[3]] == ["Part Number", "Description", "Qty", "Material", "Process", "REV", "State", "File"]
    assert ws["H4"].value == "CD-000200 Rev - x4.3mf"
    assert (ws["A4"].fill.fgColor.rgb or "")[-6:] == YELLOW


def test_unknown_kind_is_rejected(tmp_path):
    p = _partner_manifest(tmp_path, kind="traveler")
    res = mp.run_package(str(p), str(tmp_path / "x.zip"))
    assert res["ok"] is False and "Unknown package kind" in res["error"]


def test_quote_stamps_every_page_with_quote_text(tmp_path):
    out = tmp_path / "q.zip"
    res = mp.run_package(str(_partner_manifest(tmp_path, quote=True)), str(out))
    assert res["ok"] is True, res
    z = zipfile.ZipFile(out)
    pdfs = ["00 Cover.pdf", "PDF/CD-000123 Rev A.pdf", "PDF/CD-000124 Rev -.pdf"]
    for name in pdfs:
        pages = PdfReader(BytesIO(z.read(name))).pages
        assert pages
        for p in pages:
            text = p.extract_text() or ""
            assert "FOR QUOTE - NOT FOR PRODUCTION" in text, name
            assert "NOT RELEASED" not in text, name
    cover = "".join(p.extract_text() or "" for p in PdfReader(BytesIO(z.read("00 Cover.pdf"))).pages)
    assert "Quote package: CD-001882" in cover
    ws = openpyxl.load_workbook(BytesIO(z.read("BOM.xlsx")))["BOM"]
    assert ws["A1"].value == "CD-001882 Acme Machining CNC Quote package"


def test_quote_false_behaves_as_before(tmp_path):
    out = tmp_path / "q.zip"
    mp.run_package(str(_partner_manifest(tmp_path, quote=False)), str(out))
    z = zipfile.ZipFile(out)
    text = "".join(p.extract_text() or "" for p in PdfReader(BytesIO(z.read("PDF/CD-000123 Rev A.pdf"))).pages)
    assert "FOR QUOTE" not in text
    assert "Partner package" in "".join(p.extract_text() or "" for p in PdfReader(BytesIO(z.read("00 Cover.pdf"))).pages)


def test_cli_package_writes_result(tmp_path):
    out = tmp_path / "p.zip"
    r = tmp_path / "r.json"
    code = cli.main(["package", "--manifest", str(_partner_manifest(tmp_path)), "--out", str(out), "--result", str(r)])
    res = json.loads(r.read_text(encoding="utf-8"))
    assert code == 0 and res["ok"] is True and res["mode"] == "package" and Path(res["zip"]) == out


def _text(data: bytes) -> str:
    return "".join(p.extract_text() or "" for p in PdfReader(BytesIO(data)).pages)


def test_quote_stamp_fits_letter_pages(tmp_path):
    for size in ((612, 792), (792, 612)):
        c = canvas.Canvas(str(tmp_path / "d.pdf"), pagesize=size)
        c.drawString(50, 50, "drawing"); c.showPage(); c.save()
        out = mp._stamp((tmp_path / "d.pdf").read_bytes(), True)
        t = _text(out)
        assert "FOR" in t and "PRODUCTION" in t and "QUOTE" in t
        page = PdfReader(BytesIO(out)).pages[0]
        assert mp._fit_size(mp.QUOTE_TEXT, size[0], size[1], 60) * 1.0 <= 60
        # rotated bbox of the stamp stays inside the page
        from reportlab.pdfbase.pdfmetrics import stringWidth
        fs = mp._fit_size(mp.QUOTE_TEXT, size[0], size[1], 60)
        half = (stringWidth(mp.QUOTE_TEXT, "Helvetica-Bold", fs) + fs) * 0.7071 / 2
        assert half <= min(size) / 2 and float(page.mediabox.width) == size[0]


def _manifest_with_files(tmp_path, files):
    m = json.loads(_partner_manifest(tmp_path).read_text(encoding="utf-8"))
    m["files"] = files
    m["rows"] = []
    p = tmp_path / "mf.json"; p.write_text(json.dumps(m), encoding="utf-8")
    return p


def test_illegal_chars_and_duplicates_reported_in_failed(tmp_path):
    a = tmp_path / "a.stp"; a.write_bytes(b"x")
    files = [{"src": str(a), "dest": "STEP/a:b.stp", "stamp": False},
             {"src": str(a), "dest": "STEP/a?b.stp", "stamp": False},
             {"src": str(a), "dest": "STEP/ok.stp", "stamp": False},
             {"src": str(a), "dest": "STEP/OK.stp", "stamp": False}]
    out = tmp_path / "p.zip"
    res = mp.run_package(str(_manifest_with_files(tmp_path, files)), str(out))
    assert res["ok"] is True and res["step"] == 1 and res["failed"] == 1 + 3
    assert [n for n in zipfile.ZipFile(out).namelist() if n.startswith("STEP")] == ["STEP/ok.stp"]


def test_unreadable_pdf_goes_to_failed(tmp_path):
    bad = tmp_path / "bad.pdf"; bad.write_bytes(b"not a pdf")
    files = [{"src": str(bad), "dest": "PDF/bad.pdf", "stamp": True}]
    out = tmp_path / "p.zip"
    res = mp.run_package(str(_manifest_with_files(tmp_path, files)), str(out))
    assert res["ok"] is True and res["pdf"] == 0 and res["stamped"] == 0 and res["failed"] == 2
    assert "PDF/bad.pdf" not in zipfile.ZipFile(out).namelist()


def test_qty_tolerates_nan_and_inf():
    assert mp._qty(float("nan")) == 0 and mp._qty(float("inf")) == 0 and mp._qty(4.0) == 4
