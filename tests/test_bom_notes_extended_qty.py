"""Notes (the part's Comments iProperty) and extended quantities (row QTY x every ancestor's)."""
from io import BytesIO

import openpyxl
import pandas as pd
from pypdf import PdfReader

import bom_purchasing as bp
import review_package as rp

COLS = bp.SHEET_COLUMNS
NOTE = "PURCHASE AS NEEDED. BUY INSTALL TOOL: McMaster 90240A320 (installation tool for M6 tapping insert)."


def _letter(name):
    return openpyxl.utils.get_column_letter(COLS.index(name) + 1)


def _nested(**part):
    """ASM -> sub-assembly x2 -> part x36, plus the same part x4 straight under the root."""
    base = {"Source": "Buy", "Units": "Each", "Vendor": "Acme", "Cost": 0.5}
    rows = [
        {"Number": "ASM", "Row Order": "1", "Item Qty": 1, "Source": "Make", "Description (Item,CO)": "top"},
        {"Number": "SUB", "Row Order": "1.1", "Item Qty": 2, "Source": "Make", "Description (Item,CO)": "sub"},
        {**base, "Number": "CD-001162", "Row Order": "1.1.1", "Item Qty": 36, "Description (Item,CO)": "insert", **part},
        {**base, "Number": "CD-001162", "Row Order": "1.2", "Item Qty": 4, "Description (Item,CO)": "insert", **part},
    ]
    return pd.DataFrame(rows)


def _generate(tmp_path, df, monkeypatch):
    ref = pd.DataFrame({"Number": ["NONE"], "Vendor": ["V"]})

    def fake(d, reference_path="", info=None):
        if info is not None:
            info["source"] = "excel"
            info["signin_needed"] = False
        d, matched, total = bp.lookup_purchased_data(d, ref)
        return d, matched, total, [], []

    monkeypatch.setattr(bp, "_enrich_with_reference", fake)
    res = bp.generate_from_dataframe(df, "ASM", str(tmp_path))
    return openpyxl.load_workbook(res["output_path"])


def _by_vendor_rows(wb):
    ws = wb["By Vendor"]
    header = [c.value for c in ws[2]]
    out = []
    for r in ws.iter_rows(min_row=3, values_only=True):
        if r[0] is None or r[1] is None:
            break
        out.append(dict(zip(header, r)))
    return out


def test_extended_quantities_multiply_ancestors():
    df = _nested()
    assert bp.extended_quantities(df) == [1.0, 2.0, 72.0, 4.0]


def test_by_vendor_total_qty_uses_extended_quantity(tmp_path, monkeypatch):
    wb = _generate(tmp_path, _nested(), monkeypatch)
    rows = _by_vendor_rows(wb)
    assert len(rows) == 1
    assert rows[0]["Total Qty"] == 76  # 36 x 2 under the sub-assembly + 4 under the root
    assert rows[0]["Line Total"] == "=G3*F3"


def test_purchasing_tab_money_math_uses_total_qty(tmp_path, monkeypatch):
    ws = _generate(tmp_path, _nested(), monkeypatch)["Purchasing"]
    tq, cost = _letter("Total Qty"), _letter("Cost Per")
    first = bp.HDR_ROW + 1
    sub = first + 1
    part = first + 2
    assert ws[f"{tq}{first}"].value == f"=IF(ISNUMBER({_letter('Item Qty')}{first}),{_letter('Item Qty')}{first},1)"
    assert ws[f"{tq}{part}"].value.endswith(f"*{tq}{sub}")
    assert ws[f"{_letter('Sub Total')}{part}"].value == f"=({cost}{part}*{tq}{part})+{_letter('Shipping')}{part}+{_letter('Tax/Tariff')}{part}"
    assert ws[f"{_letter('Tax/Tariff')}{part}"].value == f"={cost}{part}*{tq}{part}*0.07"
    # The sub-assembly's Cost Per is the cost to make ONE: its children's sub totals spread over its own Total Qty.
    assert ws[f"{cost}{sub}"].value.endswith(f"/IF({tq}{sub}>0,{tq}{sub},1)")
    assert ws[f"{_letter('Sub Total')}{sub}"].value == f"={cost}{sub}*{tq}{sub}"
    # The hidden "paid" helper is the same expression as the line's money.
    paid_col = len(COLS) + 3
    assert ws.cell(row=part, column=paid_col).value == f"={cost}{part}*{tq}{part}"


def test_mcmaster_pack_rounding_uses_extended_quantity(tmp_path, monkeypatch):
    mc = {"McMaster Pack Price": 4.68, "Pack Size": 50, "McMaster Price Date": "2026-10-14",
          "Vendor": "McMASTER-CARR", "Vendor Number": "91251A054", "Cost": None}
    wb = _generate(tmp_path, _nested(**mc), monkeypatch)
    ws = wb["Purchasing"]
    first = bp.HDR_ROW + 1
    packs = _letter("Packs to Buy")
    assert ws[f"{packs}{first + 2}"].value == 2   # ceil(72 / 50), not ceil(36 / 50)
    assert ws[f"{packs}{first + 3}"].value == 1   # ceil(4 / 50)
    rows = _by_vendor_rows(wb)
    assert rows[0]["Total Qty"] == 76
    assert rows[0]["Line Total"].startswith("=ROUNDUP(F3/50,0)")


def test_notes_on_purchasing_and_by_vendor(tmp_path, monkeypatch):
    df = _nested()
    df["Notes"] = ["", "", NOTE, NOTE + " "]
    wb = _generate(tmp_path, df, monkeypatch)
    ws = wb["Purchasing"]
    assert COLS[-1] == "Notes"
    assert ws.cell(row=bp.HDR_ROW, column=len(COLS)).value == "Notes"
    assert ws.cell(row=bp.HDR_ROW + 3, column=len(COLS)).value == NOTE
    assert ws.cell(row=bp.HDR_ROW + 1, column=len(COLS)).value is None
    rows = _by_vendor_rows(wb)
    assert len(rows) == 1                      # Notes never split the vendor group
    assert rows[0]["Notes"] == NOTE            # distinct notes only


def test_by_vendor_joins_distinct_notes(tmp_path, monkeypatch):
    df = _nested()
    df["Notes"] = ["", "", "first", "second"]
    rows = _by_vendor_rows(_generate(tmp_path, df, monkeypatch))
    assert len(rows) == 1 and rows[0]["Notes"] == "first; second"


def test_missing_notes_column_still_works(tmp_path, monkeypatch):
    wb = _generate(tmp_path, _nested(), monkeypatch)
    assert wb["Purchasing"].cell(row=bp.HDR_ROW + 3, column=len(COLS)).value is None
    assert _by_vendor_rows(wb)[0]["Notes"] is None


def _bom_file(path, with_notes):
    head = ["Item", "Part Number", "BOM Structure", "Unit QTY", "QTY", "Description", "REV", "State"]
    rows = [["1", "CD-001882", "Normal", "Each", "1", "top", "A", "Released"],
            ["4", "CD-001887", "Normal", "Each", "2", "sub", "A", "Released"],
            ["4.3", "CD-001162", "Purchased", "Each", "36", "insert", "A", "Released"]]
    if with_notes:
        head.append("Notes")
        for r, n in zip(rows, ["", "", NOTE]):
            r.append(n)
    path.write_text("\r\n".join("\t".join(r) for r in [head] + rows) + "\r\n", encoding="utf-8")
    return str(path)


def test_package_table_shows_total_qty_and_notes(tmp_path):
    data, _, _ = rp.bom_table_data(_bom_file(tmp_path / "bom.txt", True))
    assert data[0][3:5] == ["QTY", "Total QTY"] and data[0][-1] == "Notes"
    assert [r[3:5] for r in data[1:]] == [["1", "1"], ["2", "2"], ["36", "72"]]
    assert [r[-1] for r in data[1:]] == ["", "", NOTE]
    pdf = rp._bom_table(_bom_file(tmp_path / "bom2.txt", True))
    assert "PURCHASE AS NEEDED" in " ".join((PdfReader(BytesIO(pdf)).pages[0].extract_text() or "").split())


def test_package_table_without_notes_column(tmp_path):
    data, _, _ = rp.bom_table_data(_bom_file(tmp_path / "bom.txt", False))
    assert "Notes" not in data[0] and "Total QTY" in data[0]
    assert rp._bom_table(_bom_file(tmp_path / "bom2.txt", False))
