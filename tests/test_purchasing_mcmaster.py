"""Live McMaster-Carr prices from the BOM file: pack price, Pack Size, Packs to Buy."""
import openpyxl
import pandas as pd
import pytest

import bom_purchasing as bp
import vault_tools_cli as cli

COLS = bp.SHEET_COLUMNS


def _col(name):
    return COLS.index(name) + 1


def _letter(name):
    return openpyxl.utils.get_column_letter(_col(name))


def _bom(rows):
    base = {"Number": "", "Row Order": "1", "Item Qty": 2, "Units": "Each",
            "Description (Item,CO)": "x", "Source": "Buy"}
    return pd.DataFrame([{**base, "File Name": r.get("Number", "") + ".ipt", **r} for r in rows])


def _fake_enrich(df, ref, info):
    if info is not None:
        info["source"] = "excel"
        info["signin_needed"] = False
    df, matched, total = bp.lookup_purchased_data(df, ref)
    return df, matched, total, [], []


def _generate(tmp_path, rows, monkeypatch, ref=None, rate=None):
    kw = {} if rate is None else {"sales_tax_rate": rate}
    if ref is None:
        ref = pd.DataFrame({"Number": ["NONE"], "Vendor": ["V"]})
    monkeypatch.setattr(bp, "_enrich_with_reference",
                        lambda df, reference_path="", info=None: _fake_enrich(df, ref, info))
    return bp.generate_from_dataframe(_bom(rows), "ASM", str(tmp_path), **kw)


def _sheet(tmp_path, rows, monkeypatch, ref=None, rate=None):
    res = _generate(tmp_path, rows, monkeypatch, ref, rate)
    return openpyxl.load_workbook(res["output_path"])["Purchasing"]


def _row(ws, n):
    return {name: ws.cell(row=bp.HDR_ROW + n, column=_col(name)) for name in COLS}


def _mc(number="P1", qty=2, price=4.68, size=50, **extra):
    r = {"Number": number, "Item Qty": qty, "Vendor": "McMASTER-CARR", "Vendor Number": "91251A054",
         "McMaster Pack Price": price, "Pack Size": size, "McMaster Price Date": "2026-10-14"}
    r.update(extra)
    return r


def _paid(ri):
    """Packs to Buy x pack price, where the pack price is Cost Per x Pack Size."""
    return f"{_letter('Packs to Buy')}{ri}*{_letter('Cost Per')}{ri}*{_letter('Pack Size')}{ri}"


# ---- columns ---------------------------------------------------------------

def test_pack_columns_follow_price_source():
    assert COLS.index("Pack Size") == COLS.index("Price source") + 1
    assert COLS.index("Packs to Buy") == COLS.index("Pack Size") + 1


# ---- helpers ---------------------------------------------------------------

@pytest.mark.parametrize("value,expected", [
    (50, 50), ("50", 50), (" 12 ", 12), (1, 1), (None, 1), ("", 1), (0, 1), (-3, 1), ("junk", 1),
    (float("nan"), 1), (2.0, 2),
])
def test_pack_size_of(value, expected):
    assert bp.pack_size_of(value) == expected


@pytest.mark.parametrize("qty,size,expected", [
    (2, 50, 1), (50, 50, 1), (51, 50, 2), (100, 50, 2), (101, 50, 3), (1, 1, 1), (7, 1, 7),
    (0, 50, 0), ("12", 5, 3), (2.5, 1, 3), ("-", 50, None), (None, 50, None),
])
def test_packs_to_buy_rounds_up(qty, size, expected):
    assert bp.packs_to_buy(qty, size) == expected


# ---- price precedence ------------------------------------------------------

def test_mcmaster_price_per_piece_and_source(tmp_path, monkeypatch):
    r = _row(_sheet(tmp_path, [_mc(qty=2, price=4.68, size=50)], monkeypatch), 1)
    assert r["Cost Per"].value == pytest.approx(4.68 / 50)
    assert r["Price source"].value == "McMaster"
    assert r["Pack Size"].value == 50
    assert r["Packs to Buy"].value == 1


def test_mcmaster_piece_price_shows_four_decimals(tmp_path, monkeypatch):
    ws = _sheet(tmp_path, [_mc(), {"Number": "P2", "Cost": 4.5}], monkeypatch)
    rows = {r["Price source"].value: r for r in (_row(ws, 1), _row(ws, 2))}
    assert rows["McMaster"]["Cost Per"].number_format == '"$"#,##0.00##'
    assert rows["File"]["Cost Per"].number_format == '"$"#,##0.00'


def test_mcmaster_wins_over_file_and_list(tmp_path, monkeypatch):
    ref = pd.DataFrame({"Number": ["P1"], "Vendor": ["V"], "Cost Per": [9.0]})
    r = _row(_sheet(tmp_path, [_mc(price=10.0, size=4, Cost=3.0)], monkeypatch, ref), 1)
    assert (r["Cost Per"].value, r["Price source"].value) == (2.5, "McMaster")


def test_blank_pack_size_defaults_to_one(tmp_path, monkeypatch):
    r = _row(_sheet(tmp_path, [_mc(qty=3, price=1.25, size=None)], monkeypatch), 1)
    assert r["Cost Per"].value == 1.25
    assert (r["Pack Size"].value, r["Packs to Buy"].value) == (1, 3)


@pytest.mark.parametrize("price", [0, None, "", "n/a", -1])
def test_no_live_price_falls_back_to_file(tmp_path, monkeypatch, price):
    r = _row(_sheet(tmp_path, [_mc(price=price, Cost=3.0)], monkeypatch), 1)
    assert (r["Cost Per"].value, r["Price source"].value) == (3.0, "File")
    assert r["Pack Size"].value is None and r["Packs to Buy"].value is None


def test_no_live_price_falls_back_to_list(tmp_path, monkeypatch):
    ref = pd.DataFrame({"Number": ["P1"], "Vendor": ["V"], "Cost Per": [9.0]})
    r = _row(_sheet(tmp_path, [_mc(price=None)], monkeypatch, ref), 1)
    assert (r["Cost Per"].value, r["Price source"].value) == (9.0, "List")


def test_non_mcmaster_rows_unchanged(tmp_path, monkeypatch):
    ws = _sheet(tmp_path, [{"Number": "P2", "Cost": 4.5, "Vendor": "Acme"}, _mc(number="P1")], monkeypatch)
    plain = next(r for r in (_row(ws, 1), _row(ws, 2)) if r["Price source"].value == "File")
    ri = plain["Cost Per"].row
    c, q, s, t = _letter("Cost Per"), _letter("Item Qty"), _letter("Shipping"), _letter("Tax/Tariff")
    assert plain["Cost Per"].value == 4.5
    assert plain["Pack Size"].value is None and plain["Packs to Buy"].value is None
    assert plain["Tax/Tariff"].value == f"={c}{ri}*{q}{ri}*0.07"
    assert plain["Sub Total"].value == f"=({c}{ri}*{q}{ri})+{s}{ri}+{t}{ri}"


def test_bom_without_mcmaster_columns(tmp_path, monkeypatch):
    r = _row(_sheet(tmp_path, [{"Number": "P1", "Cost": 2.0}], monkeypatch), 1)
    assert r["Price source"].value == "File"
    assert r["Pack Size"].value is None and r["Packs to Buy"].value is None


# ---- totals formulas -------------------------------------------------------

def test_mcmaster_sub_total_is_packs_times_pack_price_plus_tax(tmp_path, monkeypatch):
    r = _row(_sheet(tmp_path, [_mc(qty=120, price=4.68, size=50)], monkeypatch), 1)
    ri = r["Sub Total"].row
    assert r["Packs to Buy"].value == 3
    assert r["Tax/Tariff"].value == f"={_paid(ri)}*0.07"
    assert r["Sub Total"].value == f"=({_paid(ri)})+{_letter('Shipping')}{ri}+{_letter('Tax/Tariff')}{ri}"
    # What the formulas evaluate to: 3 packs at $4.68, plus 7% tax.
    paid = r["Packs to Buy"].value * r["Cost Per"].value * r["Pack Size"].value
    assert paid == pytest.approx(3 * 4.68)
    assert paid * 1.07 == pytest.approx(15.0228)


def test_mcmaster_tax_uses_configured_rate(tmp_path, monkeypatch):
    r = _row(_sheet(tmp_path, [_mc()], monkeypatch, rate=0.0825), 1)
    assert r["Tax/Tariff"].value == f"={_paid(r['Tax/Tariff'].row)}*0.0825"


def test_tariffed_mcmaster_row_keeps_hand_entered_tax(tmp_path, monkeypatch):
    r = _row(_sheet(tmp_path, [_mc(**{"HS/HTS Code": "7318.15", "Tax/Tariff": 1.5})], monkeypatch), 1)
    ri = r["Sub Total"].row
    assert r["Tax/Tariff"].value == 1.5
    assert r["Sub Total"].value == f"=({_paid(ri)})+{_letter('Shipping')}{ri}+{_letter('Tax/Tariff')}{ri}"


def test_tariffed_mcmaster_row_without_tax_is_flagged(tmp_path, monkeypatch):
    r = _row(_sheet(tmp_path, [_mc(**{"HS/HTS Code": "7318.15"})], monkeypatch), 1)
    assert r["Tax/Tariff"].value is None
    assert r["Tax/Tariff"].fill.fgColor.rgb.endswith(bp.TARIFF_FILL)


def test_assembly_rolls_up_pack_sub_totals(tmp_path, monkeypatch):
    rows = [
        {"Number": "ASM1", "Row Order": "1", "Item Qty": 1, "Source": "Make"},
        {**_mc(number="P1", qty=60, price=4.68, size=50), "Row Order": "1.1"},
    ]
    ws = _sheet(tmp_path, rows, monkeypatch)
    asm, part = _row(ws, 1), _row(ws, 2)
    assert asm["Pack Size"].value is None and asm["Packs to Buy"].value is None
    assert asm["Cost Per"].value == f"=SUM({_letter('Sub Total')}{part['Sub Total'].row})"
    assert part["Packs to Buy"].value == 2


def test_info_counts_mcmaster_rows(tmp_path, monkeypatch):
    res = _generate(tmp_path, [_mc(number="P1"), _mc(number="P2", price=None), {"Number": "P3"}], monkeypatch)
    assert res["mcmaster_priced"] == 1


def test_sheet_reads_back_pack_columns(tmp_path, monkeypatch):
    res = _generate(tmp_path, [_mc(qty=120)], monkeypatch)
    df, _, err = bp.read_purchasing_sheet(res["output_path"])
    assert err is None
    assert int(df.loc[0, "Pack Size"]) == 50 and int(df.loc[0, "Packs to Buy"]) == 3


# ---- the BOM file the add-in writes ---------------------------------------

BOM_HEADERS = ["Item", "Part Number", "BOM Structure", "Unit QTY", "QTY", "Description", "REV", "State", "Material",
               "Process", "File Name", "Vendor", "Vendor Number", "Cost", "HS/HTS Code",
               "McMaster Pack Price", "Pack Size", "McMaster Price Date"]


def test_cli_reads_mcmaster_columns_from_bom_file(tmp_path, monkeypatch):
    monkeypatch.setattr(bp, "_enrich_with_reference",
                        lambda df, reference_path="", info=None: (df, 0, 0, [], []))
    monkeypatch.setattr(bp, "_fill_state_from_vault", lambda df: [])
    bom = tmp_path / "bom.txt"
    rows = [
        ["1", "CD-000001", "Purchased", "Each", "120", "Socket Head Screw", "", "Released", "Steel", "", "CD-000001.ipt",
         "McMASTER-CARR", "91251A054", "0.5", "", "4.68", "50", "2026-10-14"],
        ["2", "CD-000002", "Purchased", "Each", "3", "Bracket", "", "Released", "Steel", "", "CD-000002.ipt",
         "Acme", "B-1", "2.25", "", "", "", ""],
    ]
    bom.write_text("\r\n".join("\t".join(r) for r in [BOM_HEADERS] + rows) + "\r\n", encoding="utf-8")
    out = cli.run_purchasing(str(bom), "ASM", str(tmp_path), "ASM.xlsx", source="Model state: x")
    assert out["ok"] and out["mcmaster_priced"] == 1
    ws = openpyxl.load_workbook(out["workbook"])["Purchasing"]
    by_vendor = {r["Vendor"].value: r for r in (_row(ws, 1), _row(ws, 2))}
    mc, plain = by_vendor["McMASTER-CARR"], by_vendor["Acme"]
    assert (mc["Price source"].value, mc["Pack Size"].value, mc["Packs to Buy"].value) == ("McMaster", 50, 3)
    assert mc["Cost Per"].value == pytest.approx(0.0936)
    assert (plain["Price source"].value, plain["Cost Per"].value, plain["Pack Size"].value) == ("File", 2.25, None)
