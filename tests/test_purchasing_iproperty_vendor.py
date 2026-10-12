"""Round 3: Vendor / Vendor Number from the BOM file fill what the list leaves blank,
and the engine reports which reference source it used."""
import os
import sys

import openpyxl
import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

import bom_purchasing as bp  # noqa: E402
import purchasing_reference as pref  # noqa: E402
import vault_tools_cli as cli  # noqa: E402

HEAD = "Item\tPart Number\tBOM Structure\tUnit QTY\tQTY\tDescription\tREV\tMaterial\tFile Name"


def _ref(tmp_path, monkeypatch):
    ref = tmp_path / "ref.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "purchased parts"
    ws.append(["Number", "Vendor", "Vendor Number", "Cost Per"])
    ws.append(["CD-000002", "ListVendor", "L-2", 3.5])
    ws.append(["CD-000004", "ListVendor4", None, 1.0])
    wb.save(ref)
    monkeypatch.setattr(pref, "resolve_reference_config", lambda override=None: {
        "source": "excel", "mslist": {}, "column_map": {}})
    monkeypatch.setattr(bp, "find_purchased_items_file", lambda: str(ref))


def _bom(tmp_path, with_vendor):
    rows = [
        ("1", "SF-1", "Purchased", "4", "bolt", "CD-000001.ipt", "iPropVendor", "IP-1"),
        ("2", "SF-2", "Purchased", "2", "nut", "CD-000002.ipt", "iPropOther", "IP-2"),
        ("3", "SF-3", "Purchased", "2", "washer", "CD-000003.ipt", "", ""),
        ("4", "SF-4", "Purchased", "1", "pin", "CD-000004.ipt", "iProp4", "IP-4"),
    ]
    head = HEAD + ("\tVendor\tVendor Number" if with_vendor else "")
    lines = [head]
    for i, pn, st, q, d, f, v, vn in rows:
        line = f"{i}\t{pn}\t{st}\tEach\t{q}\t{d}\t\t\t{f}"
        if with_vendor:
            line += f"\t{v}\t{vn}"
        lines.append(line)
    p = tmp_path / "bom.txt"
    p.write_text("\r\n".join(lines) + "\r\n", encoding="utf-8")
    return p


def _sheet(path):
    ws = openpyxl.load_workbook(path)["Purchasing"] if "Purchasing" in openpyxl.load_workbook(path).sheetnames \
        else openpyxl.load_workbook(path).worksheets[0]
    out = {}
    hdr = None
    for row in ws.iter_rows(values_only=True):
        if row and "Vendor" in row and "Vendor Number" in row:
            hdr = {c: i for i, c in enumerate(row) if c}
            continue
        if hdr and row and row[hdr.get("Description (Item,CO)", hdr.get("Description", 0))]:
            out[row[hdr.get("Description (Item,CO)", hdr.get("Description", 0))]] = (
                row[hdr["Vendor"]], row[hdr["Vendor Number"]])
    return out


def test_bom_vendor_fills_blanks_and_reference_wins(tmp_path, monkeypatch):
    _ref(tmp_path, monkeypatch)
    bom = _bom(tmp_path, True)
    res = bp.generate_from_file(str(bom), "A", str(tmp_path))
    assert not res.get("error"), res
    rows = _sheet(res["output_path"])
    assert rows["bolt"] == ("iPropVendor", "IP-1")        # not on list: iProperty values
    assert rows["nut"] == ("ListVendor", "L-2")           # list wins over iProperty
    assert rows["pin"] == ("ListVendor4", "IP-4")         # list vendor, blank number filled
    assert rows["washer"] == (None, None) or rows["washer"] == ("", "")
    # unmatched keeps meaning "not on the purchasing list"
    assert sorted(res["unmatched_parts"]) == ["SF-1", "SF-3"]
    assert res["reference_source"] == "excel"


def test_bom_without_vendor_columns_unchanged(tmp_path, monkeypatch):
    _ref(tmp_path, monkeypatch)
    res = bp.generate_from_file(str(_bom(tmp_path, False)), "A", str(tmp_path))
    rows = _sheet(res["output_path"])
    assert rows["bolt"] == (None, None) or rows["bolt"] == ("", "")
    assert rows["nut"] == ("ListVendor", "L-2")
    assert sorted(res["unmatched_parts"]) == ["SF-1", "SF-3"]


def _info(monkeypatch, cfg, **patches):
    monkeypatch.setattr(pref, "resolve_reference_config", lambda override=None: cfg)
    for k, v in patches.items():
        monkeypatch.setattr(pref, k, v) if hasattr(pref, k) else monkeypatch.setattr(bp, k, v)
    import pandas as pd
    df = pd.DataFrame({"Number": ["X"], "Filename": ["CD-000002.ipt"], "Source": ["Buy"]})
    info = {}
    bp._enrich_with_reference(df, info=info)
    return info


MSCFG = {"source": "auto", "column_map": {},
         "mslist": {"tenant_id": "T", "client_id": "C", "site_hostname": "h",
                    "site_path": "/s", "list_name": "L", "list_id": ""}}


def _ref_df():
    import pandas as pd
    return pd.DataFrame({"Number": ["CD-000002"], "Vendor": ["V"], "Cost Per": [1.0]})


def test_info_reports_mslist(monkeypatch):
    info = _info(monkeypatch, MSCFG, load_mslist_dataframe=lambda *a, **k: _ref_df())
    assert info == {"source": "mslist", "signin_needed": False}


def test_info_reports_excel_when_not_signed_in(monkeypatch, tmp_path):
    def boom(*a, **k):
        raise RuntimeError("not signed in")
    p = tmp_path / "r.xlsx"
    _ref_df().to_excel(p, index=False)
    info = _info(monkeypatch, MSCFG, load_mslist_dataframe=boom,
                 has_cached_login=lambda: False, find_purchased_items_file=lambda: str(p))
    assert info == {"source": "excel", "signin_needed": True}


def test_info_reports_none(monkeypatch):
    info = _info(monkeypatch, {"source": "excel", "mslist": {}, "column_map": {}},
                 find_purchased_items_file=lambda: "")
    assert info == {"source": "none", "signin_needed": False}


@pytest.mark.parametrize("source,signin,text", [
    ("mslist", False, "Prices from the purchasing Microsoft List."),
    ("excel", True, "Not signed in to the purchasing Microsoft List; used purchased items.xlsx. "
                    "Sign in once with the BOM Purchasing tool's Sign in button."),
    ("excel", False, "Prices from purchased items.xlsx."),
    ("none", False, "No purchasing list found; vendor and cost columns are blank."),
])
def test_cli_result_carries_reference_and_note(tmp_path, monkeypatch, source, signin, text):
    bom = _bom(tmp_path, True)
    real = bp.generate_from_file

    def fake(bom_path, asm, out, **k):
        res = real(bom_path, asm, out, **k)
        res["reference_source"] = source
        res["reference_signin_needed"] = signin
        return res
    _ref(tmp_path, monkeypatch)
    monkeypatch.setattr(bp, "generate_from_file", fake)
    out = cli.run_purchasing(str(bom), "A", str(tmp_path), "A Purchasing.xlsx")
    assert out["ok"] is True, out
    assert out["reference"] == source
    assert out["reference_note"] == text


def test_token_cache_is_per_user_outside_the_exe(monkeypatch):
    monkeypatch.setenv("LOCALAPPDATA", r"C:\Users\x\AppData\Local")
    p = pref._token_cache_path()
    assert p.startswith(r"C:\Users\x\AppData\Local")
    assert "_MEI" not in p
