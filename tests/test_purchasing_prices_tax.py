"""Task 8: file Cost vs list price, Price source column, shipping 0, tax rules."""
import json

import openpyxl
import pandas as pd
import pytest

import bom_purchasing as bp
import vault_tools_cli as cli

COLS = bp.SHEET_COLUMNS


def _col(name):
    return COLS.index(name) + 1


def _bom(rows):
    base = {"Number": "", "File Name": "", "Row Order": "1", "Item Qty": 2, "Units": "Each",
            "Description (Item,CO)": "x", "Source": "Buy"}
    return pd.DataFrame([{**base, "File Name": r.get("Number", "") + ".ipt", **r} for r in rows])


def _fake_enrich(df, ref, info):
    if info is not None:
        info["source"] = "excel"
        info["signin_needed"] = False
    df, matched, total = bp.lookup_purchased_data(df, ref)
    return df, matched, total, [], []


def _sheet(tmp_path, rows, monkeypatch, ref=None, rate=None):
    raw = _bom(rows)
    kw = {} if rate is None else {"sales_tax_rate": rate}
    if ref is None:
        ref = pd.DataFrame({"Number": ["NONE"], "Vendor": ["V"]})
    monkeypatch.setattr(bp, "_enrich_with_reference",
                        lambda df, reference_path="", info=None: _fake_enrich(df, ref, info))
    res = bp.generate_from_dataframe(raw, "ASM", str(tmp_path), **kw)
    return openpyxl.load_workbook(res["output_path"])["Purchasing"]


def _ref(**cols):
    return pd.DataFrame(cols)


def _row(ws, n):
    return {name: ws.cell(row=bp.HDR_ROW + n, column=_col(name)) for name in COLS}


def test_price_source_column_follows_cost_per():
    assert COLS.index("Price source") == COLS.index("Cost Per") + 1


def test_file_cost_wins_over_list(tmp_path, monkeypatch):
    ref = _ref(Number=["P1"], Vendor=["V"], **{"Cost Per": [9.0]})
    r = _row(_sheet(tmp_path, [{"Number": "P1", "Cost": 4.5}], monkeypatch, ref), 1)
    assert r["Cost Per"].value == 4.5
    assert r["Price source"].value == "File"


def test_zero_file_cost_falls_back_to_list(tmp_path, monkeypatch):
    ref = _ref(Number=["P1"], Vendor=["V"], **{"Cost Per": [9.0]})
    r = _row(_sheet(tmp_path, [{"Number": "P1", "Cost": 0}], monkeypatch, ref), 1)
    assert r["Cost Per"].value == 9.0
    assert r["Price source"].value == "List"


def test_blank_file_cost_uses_list(tmp_path, monkeypatch):
    ref = _ref(Number=["P1"], Vendor=["V"], **{"Cost Per": [9.0]})
    r = _row(_sheet(tmp_path, [{"Number": "P1", "Cost": None}], monkeypatch, ref), 1)
    assert r["Price source"].value == "List"


def test_file_cost_without_a_list_match(tmp_path, monkeypatch):
    ref = _ref(Number=["OTHER"], Vendor=["V"], **{"Cost Per": [9.0]})
    r = _row(_sheet(tmp_path, [{"Number": "P1", "Cost": 3.25}], monkeypatch, ref), 1)
    assert (r["Cost Per"].value, r["Price source"].value) == (3.25, "File")


def test_neither_price_leaves_both_blank(tmp_path, monkeypatch):
    ref = _ref(Number=["OTHER"], Vendor=["V"], **{"Cost Per": [9.0]})
    r = _row(_sheet(tmp_path, [{"Number": "P1"}], monkeypatch, ref), 1)
    assert r["Cost Per"].value is None
    assert r["Price source"].value is None


def test_no_cost_column_mcp_style_list_or_blank(tmp_path, monkeypatch):
    ref = _ref(Number=["P1"], Vendor=["V"], **{"Cost Per": [9.0]})
    ws = _sheet(tmp_path, [{"Number": "P1"}, {"Number": "P2", "Row Order": "2"}],
                monkeypatch, ref)
    assert _row(ws, 1)["Price source"].value == "List"
    assert _row(ws, 2)["Price source"].value is None


def test_shipping_is_zero_and_untariffed_tax_is_the_default_formula(tmp_path, monkeypatch):
    ref = _ref(Number=["P1"], Vendor=["V"],
               **{"Cost Per": [9.0], "Shipping": [5.0], "Tax/Tariff": [99.0]})
    r = _row(_sheet(tmp_path, [{"Number": "P1"}], monkeypatch, ref), 1)
    cost, qty = r["Cost Per"].coordinate, r["Item Qty"].coordinate
    assert r["Shipping"].value == 0
    assert r["Tax/Tariff"].value == f"={cost}*{qty}*0.07"


def test_custom_sales_tax_rate(tmp_path, monkeypatch):
    ws = _sheet(tmp_path, [{"Number": "P1", "Cost": 1.0}], monkeypatch, rate=0.0825)
    assert _row(ws, 1)["Tax/Tariff"].value.endswith("*0.0825")


def test_tariffed_row_uses_entered_value_never_the_formula(tmp_path, monkeypatch):
    ref = _ref(Number=["P1"], Vendor=["V"],
               **{"Cost Per": [9.0], "HS/HTS Code": ["7318.15"], "Tax/Tariff": [12.5]})
    r = _row(_sheet(tmp_path, [{"Number": "P1"}], monkeypatch, ref), 1)
    assert r["Tax/Tariff"].value == 12.5


def test_tariffed_row_without_value_is_orange_with_comment(tmp_path, monkeypatch):
    ref = _ref(Number=["P1"], Vendor=["V"], **{"Cost Per": [9.0], "HS/HTS Code": ["7318.15"]})
    cell = _row(_sheet(tmp_path, [{"Number": "P1"}], monkeypatch, ref), 1)["Tax/Tariff"]
    assert cell.value is None
    assert cell.fill.fgColor.rgb.endswith(bp.TARIFF_FILL)
    assert cell.comment.text == "Tariff part: enter tax/tariff"


def test_hs_code_from_the_bom_file_makes_a_tariff_row(tmp_path, monkeypatch):
    ref = _ref(Number=["OTHER"], Vendor=["V"], **{"Cost Per": [9.0]})
    r = _row(_sheet(tmp_path, [{"Number": "P1", "HS/HTS Code": "8481.80"}], monkeypatch, ref), 1)
    assert r["HS/HTS Code"].value == "8481.80"
    assert r["Tax/Tariff"].value is None
    assert r["Tax/Tariff"].comment is not None


def test_list_hs_code_wins_over_file(tmp_path, monkeypatch):
    ref = _ref(Number=["P1"], Vendor=["V"], **{"HS/HTS Code": ["1111"]})
    r = _row(_sheet(tmp_path, [{"Number": "P1", "HS/HTS Code": "2222"}], monkeypatch, ref), 1)
    assert r["HS/HTS Code"].value == "1111"


def test_file_reader_keeps_hs_code_as_text(tmp_path):
    p = tmp_path / "b.csv"
    p.write_text("Number,HS/HTS Code\nA,7318\n", encoding="utf-8")
    assert bp.read_bom_file(str(p))["HS/HTS Code"].iloc[0] == "7318"


def test_assembly_rows_keep_rollup_and_blank_price_source(tmp_path, monkeypatch):
    rows = [{"Number": "ASM", "Row Order": "1", "Source": "Make", "Item Qty": 1},
            {"Number": "P1", "Row Order": "1.1", "Cost": 2.0}]
    top = _row(_sheet(tmp_path, rows, monkeypatch), 1)
    assert str(top["Cost Per"].value).startswith("=SUM(")
    assert top["Price source"].value is None


def test_cli_passes_the_sales_tax_rate(tmp_path, monkeypatch):
    bom = tmp_path / "b.csv"
    bom.write_text("Number,Item Qty,Source,Row Order,Cost\nP1,1,Buy,1,10\n", encoding="utf-8")
    seen = {}
    real = bp.generate_from_file

    def spy(*a, **k):
        seen.update(k)
        return real(*a, **k)
    monkeypatch.setattr(bp, "generate_from_file", spy)
    monkeypatch.setattr(bp, "_enrich_with_reference",
                        lambda df, reference_path="", info=None: (df, 0, 0, [], []))
    res = tmp_path / "r.json"
    code = cli.main(["purchasing", "--bom", str(bom), "--assembly", "A", "--out", str(tmp_path),
                     "--name", "A.xlsx", "--result", str(res), "--sales-tax-rate", "0.05"])
    assert code == 0, json.loads(res.read_text())
    assert seen["sales_tax_rate"] == 0.05


def test_cli_default_rate_is_seven_percent():
    args = cli._parser().parse_args(["purchasing", "--bom", "b", "--assembly", "a", "--out", "o",
                                     "--name", "n", "--result", "r"])
    assert args.sales_tax_rate == pytest.approx(0.07)


def test_text_prices_with_dollar_signs_are_read(tmp_path, monkeypatch):
    ref = _ref(Number=["P1"], Vendor=["V"],
               **{"Cost Per": ["$9.00"], "HS/HTS Code": ["7318"], "Tax/Tariff": ["$12.50"]})
    r = _row(_sheet(tmp_path, [{"Number": "P1", "Cost": "$4.50"}, ], monkeypatch, ref), 1)
    assert (r["Cost Per"].value, r["Price source"].value) == (4.5, "File")
    assert r["Tax/Tariff"].value == 12.5


@pytest.mark.parametrize("bad", ["-0.1", "7", "abc"])
def test_cli_rejects_a_bad_sales_tax_rate(bad):
    with pytest.raises(SystemExit):
        cli._parser().parse_args(["purchasing", "--bom", "b", "--assembly", "a", "--out", "o",
                                  "--name", "n", "--result", "r", "--sales-tax-rate", bad])


def test_addin_style_bom_file_with_cost_and_hs_code_end_to_end(tmp_path, monkeypatch):
    hdr = ("Item\tPart Number\tBOM Structure\tUnit QTY\tQTY\tDescription\tREV\tState\tMaterial\tProcess\t"
           "File Name\tVendor\tVendor Number\tCost\tHS/HTS Code")
    rows = [
        "1\tCD-990020\tPurchased\tEach\t2\tBolt\t\tReleased\tSteel\t\tCD-990020.ipt\tMcMaster\t91251A1\t4.5\t7318.15",
        "2\tCD-990021\tPurchased\tEach\t1\tNut\t\tReleased\tSteel\t\tCD-990021.ipt\tMcMaster\t91251A2\t\t",
    ]
    bom = tmp_path / "bom.txt"
    bom.write_text(hdr + "\r\n" + "\r\n".join(rows) + "\r\n", encoding="utf-8")
    ref = pd.DataFrame({"Number": ["CD-990021"], "Vendor": ["V"], "Cost Per": [9.0]})
    monkeypatch.setattr(bp, "_enrich_with_reference",
                        lambda df, reference_path="", info=None: _fake_enrich(df, ref, info))
    res = bp.generate_from_file(str(bom), "ASM", str(tmp_path))
    assert not res.get("error"), res
    ws = openpyxl.load_workbook(res["output_path"])["Purchasing"]
    r1, r2 = _row(ws, 1), _row(ws, 2)
    assert (r1["Cost Per"].value, r1["Price source"].value) == (4.5, "File")
    assert r1["HS/HTS Code"].value == "7318.15"
    assert r2["Cost Per"].value == 9.0
    assert r2["Price source"].value != "File"
