import json
import shutil
from pathlib import Path

import vault_tools_cli as cli

FIXTURE = Path(__file__).parent / "fixtures" / "CD-001608-inventor-bom.txt"


def test_purchasing_writes_named_workbook_and_result(tmp_path):
    bom = tmp_path / "bom.txt"
    shutil.copy(FIXTURE, bom)
    result_path = tmp_path / "result.json"
    code = cli.main(["purchasing", "--bom", str(bom), "--assembly", "CD-001608",
                     "--out", str(tmp_path), "--name", "CD-001608 Purchasing.xlsx",
                     "--result", str(result_path)])
    result = json.loads(result_path.read_text(encoding="utf-8"))
    assert code == 0, result
    assert result["ok"] is True
    assert Path(result["workbook"]) == tmp_path / "CD-001608 Purchasing.xlsx"
    assert Path(result["workbook"]).is_file()
    assert not (tmp_path / "CD-001608-PurchasingExport.xlsx").exists()
    assert result["lines"] > 0
    assert isinstance(result["unmatched"], list) and isinstance(result["warnings"], list)


def test_purchasing_replaces_existing_workbook(tmp_path):
    bom = tmp_path / "bom.txt"
    shutil.copy(FIXTURE, bom)
    (tmp_path / "A Purchasing.xlsx").write_bytes(b"old")
    cli.main(["purchasing", "--bom", str(bom), "--assembly", "A", "--out", str(tmp_path),
              "--name", "A Purchasing.xlsx", "--result", str(tmp_path / "r.json")])
    assert (tmp_path / "A Purchasing.xlsx").read_bytes()[:2] == b"PK"


def test_purchasing_reads_quoted_description(tmp_path):
    bom = tmp_path / "bom.txt"
    bom.write_text('Item\tPart Number\tBOM Structure\tUnit QTY\tQTY\tDescription\tREV\tMaterial\tFile Name\r\n'
                   '1\tCD-000001\tPurchased\tEach\t4\t"1/2"" bolt"\t\tSteel\tCD-000001.ipt\r\n', encoding="utf-8")
    out = cli.run_purchasing(str(bom), "X", str(tmp_path), "X Purchasing.xlsx")
    assert out["ok"] is True, out
    assert out["lines"] == 1


def test_failure_is_reported_in_result_file(tmp_path):
    result_path = tmp_path / "r.json"
    code = cli.main(["purchasing", "--bom", str(tmp_path / "missing.txt"), "--assembly", "X",
                     "--out", str(tmp_path), "--name", "X Purchasing.xlsx", "--result", str(result_path)])
    result = json.loads(result_path.read_text(encoding="utf-8"))
    assert code == 1
    assert result["ok"] is False and result["error"]


def test_unexpected_exception_still_writes_result(tmp_path, monkeypatch):
    monkeypatch.setattr(cli, "run_purchasing", lambda *a: (_ for _ in ()).throw(RuntimeError("boom")))
    result_path = tmp_path / "r.json"
    code = cli.main(["purchasing", "--bom", "x", "--assembly", "X", "--out", str(tmp_path),
                     "--name", "n.xlsx", "--result", str(result_path)])
    assert code == 1
    assert "boom" in json.loads(result_path.read_text(encoding="utf-8"))["error"]


ADDIN_HEADERS = ("Item\tPart Number\tBOM Structure\tUnit QTY\tQTY\tDescription\tREV\tState\t"
                 "Material\tFile Name\tVendor\tVendor Number")


def _addin_bom(path):
    path.write_text(ADDIN_HEADERS + "\r\n"
                    "1\tCD-001882\tNormal\tEach\t1\tTop\tA\t\t\tCD-001882.iam\t\t\r\n"
                    "2\tCD-000001\tPurchased\tEach\t4\tBolt\t\t\tSteel\tCD-000001.ipt\t\t\r\n",
                    encoding="utf-8")


def _purchasing_sheet(tmp_path, *extra):
    import openpyxl
    bom = tmp_path / "bom.txt"
    _addin_bom(bom)
    code = cli.main(["purchasing", "--bom", str(bom), "--assembly", "CD-001882", "--out", str(tmp_path),
                     "--name", "P.xlsx", "--result", str(tmp_path / "r.json"), *extra])
    assert code == 0, (tmp_path / "r.json").read_text(encoding="utf-8")
    return openpyxl.load_workbook(tmp_path / "P.xlsx")["Purchasing"]


def test_source_goes_on_the_generated_row(tmp_path):
    src = "Model state: Primary | Design view: Default | BOM view: Structured"
    ws = _purchasing_sheet(tmp_path, "--source", src)
    assert ws["A2"].value.startswith("Generated: ")
    assert ws["A2"].value.endswith(" | " + src)


def test_generated_row_unchanged_without_source(tmp_path):
    ws = _purchasing_sheet(tmp_path)
    assert "|" not in ws["A2"].value and "Model state" not in ws["A2"].value


def test_source_argument_parses():
    args = cli._parser().parse_args(["purchasing", "--bom", "b", "--assembly", "A", "--out", "o",
                                     "--name", "n", "--result", "r", "--source", "S"])
    assert args.source == "S"
    args = cli._parser().parse_args(["purchasing", "--bom", "b", "--assembly", "A", "--out", "o",
                                     "--name", "n", "--result", "r"])
    assert args.source is None


def test_name_column_is_filled_from_an_addin_bom(tmp_path):
    ws = _purchasing_sheet(tmp_path)
    headers = [c.value for c in ws[3]]
    col = headers.index("Name") + 1
    assert [ws.cell(row=r, column=col).value for r in (4, 5)] == ["CD-001882", "CD-000001"]
