"""Totals block on the Purchasing tab and the vendor summary on By Vendor."""
import re

import openpyxl
import pandas as pd

import bom_purchasing as bp

COLS = bp.SHEET_COLUMNS
FIRST = bp.HDR_ROW + 1


def _letter(name):
    return openpyxl.utils.get_column_letter(COLS.index(name) + 1)


def _build(tmp_path, rows):
    df = pd.DataFrame(rows).reset_index(drop=True)
    for c in bp.ALL_COLUMNS:
        if c not in df.columns:
            df[c] = None
    out = str(tmp_path / "t.xlsx")
    bp.build_purchasing_sheet(df, out, "ASM", sales_tax_rate=0.07)
    return openpyxl.load_workbook(out)


def _val(ws, ref):
    """Evaluate a data cell: a number, or a product / sum of cell refs."""
    v = ws[ref].value
    if isinstance(v, str) and v.startswith("="):
        expr = re.sub(r"\$?([A-Z]{1,3})\$?(\d+)",
                      lambda m: repr(_val(ws, m.group(1) + m.group(2))), v[1:])
        expr = re.sub(r"IF\(ISNUMBER\(([^,]+)\),\1,1\)", r"\1", expr)
        return eval(expr)
    return 0 if v is None else v


def _rows():
    return [
        {"Number": "ASM", "Title": "Top", "Item Qty": 1, "Category Name": "Assembly",
         "Row Order": "1"},
        {"Number": "F1", "Title": "file part", "Item Qty": 3, "Row Order": "1.1",
         "Source": "Buy", "Vendor": "Acme", "Cost Per": 10.0},
        {"Number": "L1", "Title": "list part", "Item Qty": 2, "Row Order": "1.2",
         "Source": "Buy", "Vendor": "Acme", "Cost Per": 2.5, "HS/HTS Code": "7318",
         "Tax/Tariff": 1.25},
        {"Number": "M1", "Title": "mcm", "Item Qty": 120, "Row Order": "1.3",
         "Source": "Buy", "Vendor": "McMASTER-CARR", "Cost Per": 0.1,
         "Price source": bp.PRICE_SOURCE_MCMASTER, "Pack Size": 50, "Packs to Buy": 3},
        {"Number": "N1", "Title": "no price", "Item Qty": 4, "Row Order": "1.4",
         "Source": "Buy", "Vendor": "Beta"},
    ]


def _totals(ws):
    out = {}
    for r in range(1, ws.max_row + 1):
        if ws.cell(row=r, column=1).value == "Totals":
            for k in range(1, 7):
                out[ws.cell(row=r + k, column=1).value] = ws.cell(row=r + k, column=2)
    return out


def test_totals_block_values_and_formulas(tmp_path):
    wb = _build(tmp_path, _rows())
    ws = wb["Purchasing"]
    t = _totals(ws)
    last = FIRST + 4
    flag_col = openpyxl.utils.get_column_letter(len(COLS) + 2)
    paid_col = openpyxl.utils.get_column_letter(len(COLS) + 3)
    # flags: assembly row 0, four part rows 1
    assert [ws[f"{flag_col}{r}"].value for r in range(FIRST, last + 1)] == [0, 1, 1, 1, 1]
    assert ws[f"{paid_col}{FIRST}"].value is None
    # McMaster row pays whole packs: 3 packs x 0.1 x 50
    assert _val(ws, f"{paid_col}{FIRST + 3}") == 3 * 0.1 * 50
    paid = sum(_val(ws, f"{paid_col}{r}") for r in range(FIRST, last + 1))
    assert paid == 30 + 5 + 15 + 0
    assert t["Parts subtotal"].value == f"=SUM(${paid_col}${FIRST}:${paid_col}${last})"
    assert f"${_letter('Shipping')}${FIRST}:${_letter('Shipping')}${last}" in t["Shipping"].value
    assert f"${_letter('Tax/Tariff')}${FIRST}:${_letter('Tax/Tariff')}${last}" in t["Tax / Tariff"].value
    g = t["Grand total"].value
    for k in ("Parts subtotal", "Shipping", "Tax / Tariff"):
        assert t[k].coordinate in g
    assert t["Lines"].value == f"=SUM(${flag_col}${FIRST}:${flag_col}${last})"
    cost = _letter("Cost Per")
    assert f"${cost}${FIRST}:${cost}${last}" in t["Lines without a price"].value
    # Grand total equals the sum of part-row Sub Totals (assembly Sub Total excluded)
    sub = _letter("Sub Total")
    sub_total = sum(_val(ws, f"{sub}{r}") for r in range(FIRST + 1, last + 1))
    tax = sum(_val(ws, f"{_letter('Tax/Tariff')}{r}") for r in range(FIRST + 1, last + 1))
    assert abs(paid + 0 + tax - sub_total) < 1e-9
    assert t["Parts subtotal"].number_format == '"$"#,##0.00'
    assert t["Parts subtotal"].coordinate == f"B{last + 4}"  # gap of two rows, header, then rows
    assert ws.cell(row=last + 3, column=1).value == "Totals"
    assert ws.cell(row=last + 4, column=1).font.bold


def test_unpriced_highlight_rule_and_note(tmp_path):
    ws = _build(tmp_path, _rows())["Purchasing"]
    t = _totals(ws)
    cell = t["Lines without a price"].coordinate
    rules = [r for rng, rs in ws.conditional_formatting._cf_rules.items()
             for r in rs if cell in str(rng.sqref)]
    assert rules and rules[0].formula == [f"{cell}>0"]
    assert rules[0].dxf.fill.bgColor.rgb.endswith("FFC000")
    note = ws.cell(row=t["Lines without a price"].row, column=3).value
    assert "Total is low: " in note and "lines have no price" in note
    assert cell in note


def test_assembly_rows_excluded_from_every_count(tmp_path):
    ws = _build(tmp_path, _rows())["Purchasing"]
    flag_col = openpyxl.utils.get_column_letter(len(COLS) + 2)
    assert ws[f"{flag_col}{FIRST}"].value == 0
    assert ws[f"{flag_col}{FIRST}"].parent.column_dimensions[flag_col].hidden


def test_reader_ignores_totals_block(tmp_path):
    wb = _build(tmp_path, _rows())
    path = str(tmp_path / "r.xlsx")
    wb.save(path)
    df, asm, err = bp.read_purchasing_sheet(path)
    assert err is None and len(df) == 5


def test_vendor_summary_and_grand_total(tmp_path):
    wb = _build(tmp_path, _rows())
    ws = wb["By Vendor"]
    top = next(r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=2).value == "Lines")
    vendors = []
    r = top + 1
    while ws.cell(row=r, column=1).value != "Grand total":
        vendors.append(ws.cell(row=r, column=1).value)
        r += 1
    assert vendors == ["Acme", "Beta", "McMASTER-CARR"]
    acme = top + 1
    assert ws.cell(row=acme, column=2).value.startswith("=COUNTIF(")
    assert ws.cell(row=acme, column=3).value.startswith("=SUMIF(")
    assert "Purchasing!" in ws.cell(row=acme, column=4).value
    assert ws.cell(row=acme, column=5).value == f"=C{acme}+D{acme}"
    assert ws.cell(row=r, column=3).value == f"=SUM(C{top + 1}:C{r - 1})"
    assert ws.cell(row=r + 2, column=5).value == "=Purchasing!$B$%d" % (
        _totals(wb["Purchasing"])["Grand total"].row)
    assert ws.cell(row=r + 3, column=5).value == f"=E{r + 2}-E{r}"
    assert "authoritative" in ws.cell(row=r + 4, column=1).value
    # each line total on this tab feeds the subtotal range
    assert f"$H$3:$H${top - 3}" in ws.cell(row=acme, column=3).value


def test_empty_bom_has_no_totals(tmp_path):
    wb = _build(tmp_path, [{"Number": "X"}])
    assert wb["Purchasing"] is not None


def test_block_clears_the_unmatched_note(tmp_path):
    rows = _rows()
    rows.append({"Number": "U1", "Title": "unmatched", "Item Qty": 1, "Row Order": "1.5",
                 "Source": "Buy"})
    ws = _build(tmp_path, rows)["Purchasing"]
    last = FIRST + 5
    assert str(ws.cell(row=last + 2, column=1).value).startswith(bp.UNMATCHED_NOTE_PREFIX)
    assert ws.cell(row=last + 3, column=1).value is None
    assert ws.cell(row=last + 4, column=1).value == "Totals"
