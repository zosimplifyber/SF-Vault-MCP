"""Manufacturing packages for the SF Vault Bridge ribbon (partner and 3D print).

Reads the manifest the add-in writes (SFVaultBridge.Core PackageManifest), stages
the exported files, writes BOM.xlsx or Print List.xlsx, builds the cover PDF
(partner only), stamps flagged drawings, and zips the result. Spec:
docs/superpowers/specs/2026-10-13-ribbon-manufacturing-packages-design.md
"""
from __future__ import annotations

import json
import math
import os
import re
import shutil
import tempfile
import zipfile
from datetime import date
from io import BytesIO
from pathlib import Path, PurePosixPath
from typing import Any

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from pypdf import PdfReader
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.platypus import Spacer

import bom_purchasing as bp
import review_package as rp
from pdf_watermark import apply_watermark_to_pages

QUOTE_TEXT = "FOR QUOTE - NOT FOR PRODUCTION"
QUOTE_FONT_SIZE = 60
MODE_LABELS = {"Make": "All Make parts", "CNC": "CNC (machined)", "Sheet Metal": "Sheet metal"}
PARTNER_COLUMNS = ["Item", "Part Number", "Description", "Qty", "Material", "Material Finish",
                   "Process", "REV", "State", "File(s) in package", "Vendor Number", "Note"]
PARTNER_WIDTHS = [8, 18, 44, 8, 22, 18, 16, 8, 18, 48, 18, 36]
LEFT_OUT_HEADING = "Left out: McMaster-Carr CAD (supplier terms)"
PRINT_COLUMNS = ["Part Number", "Description", "Qty", "Material", "Process", "REV", "State", "File"]
PRINT_WIDTHS = [18, 44, 8, 22, 16, 8, 18, 48]


ILLEGAL_NAME = re.compile(r'[<>:"|?*\\\x00-\x1f]')


class _BadManifest(Exception):
    pass


def _check_dest(dest: Any) -> str:
    d = str(dest or "")
    parts = d.split("/")
    if (not d or "\\" in d or d.startswith("/") or re.match(r"^[A-Za-z]:", d)
            or ".." in parts or "" in parts or "." in parts):
        raise _BadManifest(f"Bad package path: {d}.")
    return d


def _qty(v: Any) -> Any:
    try:
        f = float(v)
    except (TypeError, ValueError):
        return v
    if math.isnan(f) or math.isinf(f):
        return 0  # not a usable quantity; the row still lists the part
    return int(f) if f == int(f) else f


def _fit_size(text: str, page_w: float, page_h: float, cap: int) -> int:
    """Largest font size (<= cap) whose 45 degree text fits in 90% of the page."""
    unit = stringWidth(text, "Helvetica-Bold", 1.0)
    limit = 0.9 * min(page_w, page_h)
    return max(8, min(cap, int(limit / ((unit + 1.0) * math.sqrt(0.5)))))


def _stamp(data: bytes, quote: bool) -> bytes:
    text, cap = (QUOTE_TEXT, QUOTE_FONT_SIZE) if quote else (rp.UNRELEASED_TEXT, 120)
    groups: dict[int, list[int]] = {}
    for i, page in enumerate(PdfReader(BytesIO(data)).pages):
        size = _fit_size(text, float(page.mediabox.width), float(page.mediabox.height), cap)
        groups.setdefault(size, []).append(i)
    for size, idx in groups.items():
        data = apply_watermark_to_pages(data, text, idx, font_size=size,
                                        color=rp.WATERMARK_COLOR, opacity=rp.WATERMARK_OPACITY)
    return data


def _write_sheet(path: Path, m: dict[str, Any], rows: list[dict[str, Any]], partner: bool, quote: bool) -> None:
    wb = Workbook()
    ws = wb.active
    cols = PARTNER_COLUMNS if partner else PRINT_COLUMNS
    widths = PARTNER_WIDTHS if partner else PRINT_WIDTHS
    ws.title = "BOM" if partner else "Print List"
    n = len(cols)
    last = get_column_letter(n)
    asm = m.get("assembly", "")
    if partner:
        kind = "Quote package" if quote else "package"
        scope = str(m.get("scope") or "").strip()
        title = f"{scope} | {m.get('partner', '')} {m.get('mode', '')} {kind}" if scope else f"{asm} {m.get('partner', '')} {m.get('mode', '')} {kind}"
    else:
        title = f"{asm} 3D print list"
    ws.merge_cells(f"A1:{last}1")
    c = ws["A1"]
    c.value = title
    c.font = Font(name="Arial", size=11, bold=True, color=bp.WHITE)
    c.fill = PatternFill("solid", fgColor=bp.DARK_BLUE)
    c.alignment = Alignment(horizontal="center", vertical="center")
    for i in range(1, n + 1):
        ws.cell(row=1, column=i).fill = PatternFill("solid", fgColor=bp.DARK_BLUE)
    ws.row_dimensions[1].height = 50
    bp._add_logo(ws)
    sub = f"Generated: {date.today().strftime('%B %d, %Y')}"
    if str(m.get("source") or "").strip():
        sub += f" | {m['source']}"
    ws.merge_cells(f"A2:{last}2")
    c = ws["A2"]
    c.value = sub
    c.font = Font(name="Arial", size=9, italic=True, color=bp.DARK_GRAY)
    c.alignment = Alignment(horizontal="right", vertical="center")
    for i in range(1, n + 1):
        ws.cell(row=2, column=i).fill = PatternFill("solid", fgColor=bp.LIGHT_GRAY)
    thin = Side(style="thin", color=bp.GRAY_BDR)
    for i, name in enumerate(cols, 1):
        h = ws.cell(row=3, column=i, value=name)
        h.font = Font(name="Arial", size=10, bold=True, color=bp.WHITE)
        h.fill = PatternFill("solid", fgColor=bp.DARK_BLUE)
        h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        h.border = Border(left=thin, right=thin, bottom=thin, top=Side(style="medium", color=bp.MID_BLUE))
    for r, row in enumerate(rows):
        files = [os.path.basename(f) for f in row.get("files") or []]
        if partner:
            vals = [row.get("item", ""), row.get("part_number", ""), row.get("description", ""), _qty(row.get("qty")),
                    row.get("material", ""), row.get("finish", ""), row.get("process", ""), row.get("revision", ""),
                    row.get("state", ""), ", ".join(files), row.get("vendor_number", ""), row.get("note", "")]
            qty_i = 3
        else:
            vals = [row.get("part_number", ""), row.get("description", ""), _qty(row.get("qty")),
                    row.get("material", ""), row.get("process", ""), row.get("revision", ""),
                    row.get("state", ""), files[0] if files else ""]
            qty_i = 2
        if not bp.is_released(row.get("state")) and not row.get("note"):
            fill = bp.UNRELEASED_FILL
        else:
            fill = bp.PALE_BLUE if r % 2 == 0 else bp.WHITE
        for i, v in enumerate(vals):
            cell = ws.cell(row=4 + r, column=i + 1, value=v)
            cell.font = Font(name="Arial", size=10)
            cell.fill = PatternFill("solid", fgColor=fill)
            cell.border = bp._border()
            cell.alignment = Alignment(vertical="center", wrap_text=True,
                                       horizontal="center" if i == qty_i else "left")
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.auto_filter.ref = f"A3:{last}{3 + len(rows)}"
    ws.freeze_panes = "A4"
    wb.save(path)


def _cover(m: dict[str, Any], counts: dict[str, int], rows: int, failed: list[dict[str, str]], quote: bool) -> bytes:
    st = getSampleStyleSheet()
    N = st["Normal"]
    title = f"Quote package: {m.get('assembly', '')}" if quote else f"Partner package: {m.get('assembly', '')}"
    mode = str(m.get("mode") or "")
    scope = str(m.get("scope") or "").strip()
    story = [rp._para(title, st["Title"]),
             rp._para(m.get("description") or "", st["Heading2"]),
             rp._para(f"Partner: {m.get('partner', '')}", N),
             rp._para(f"Parts: {MODE_LABELS.get(mode, mode)}", N),
             *([rp._para(f"Scope: {scope}", N)] if scope else []),
             rp._para(f"{'Part' if scope.startswith('Part ') else 'Assembly'} {m.get('assembly', '')} Rev {m.get('revision') or '-'} "
                      f"({str(m.get('state') or '').strip() or 'unknown'})", N),
             rp._para(f"Created {m.get('created', '')} by {m.get('user', '')}", N),
             rp._para(f"{rows} parts: {counts['step']} STEP, {counts['pdf']} PDF, {counts['dxf']} DXF", N)]
    sections = [("Process not set (guessed)", [str(x) for x in m.get("guessed") or []]),
                ("Missing drawings", [str(x) for x in m.get("missing_drawings") or []]),
                ("No flat pattern", [str(x) for x in m.get("no_flat_pattern") or []]),
                ("Could not export", [f"{f.get('part', '')}: {f.get('reason', '')}" for f in failed])]
    left_out = [str(x) for x in m.get("left_out") or []]
    if left_out:
        sections.append((LEFT_OUT_HEADING, left_out))
    for head, items in sections:
        story += [Spacer(1, 12), rp._para(head, st["Heading3"])]
        story += [rp._para(x, N) for x in items] or [rp._para("None", N)]
    return rp._doc(story)


def _zip(staging: Path, out_zip: str) -> int:
    names = [p.relative_to(staging).as_posix() for p in staging.rglob("*") if p.is_file()]
    first = [n for n in ("00 Cover.pdf", "BOM.xlsx", "Print List.xlsx") if n in names]
    rest = sorted((n for n in names if n not in first), key=str.lower)
    part = out_zip + ".part"
    Path(out_zip).parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(part, "w", zipfile.ZIP_DEFLATED) as z:
        for n in first + rest:
            z.write(staging / n, n)
    os.replace(part, out_zip)
    return len(first) + len(rest)


def run_package(manifest: str, out_zip: str) -> dict[str, Any]:
    m = json.loads(Path(manifest).read_text(encoding="utf-8-sig"))
    kind = m.get("kind")
    try:
        if kind not in ("partner", "print"):
            raise _BadManifest(f"Unknown package kind: {kind}.")
        entries = [(e, _check_dest(e.get("dest"))) for e in m.get("files") or []]
    except _BadManifest as exc:
        return {"ok": False, "error": str(exc)}
    partner = kind == "partner"
    quote = bool(m.get("quote", False))
    staging = Path(tempfile.mkdtemp(prefix="sfpkg-"))
    try:
        failed: list[dict[str, str]] = []
        dropped: set[str] = set()
        written: list[str] = []
        stamped = 0
        seen: set[str] = set()
        for e, dest in entries:
            target = staging / dest
            stem = PurePosixPath(dest).stem
            if any(ILLEGAL_NAME.search(seg) for seg in dest.split("/")):
                failed.append({"part": stem, "reason": "illegal characters in file name"})
                dropped.add(dest)
                continue
            if dest.lower() in seen:
                failed.append({"part": stem, "reason": "duplicate file name in package"})
                continue
            seen.add(dest.lower())
            try:
                data = Path(e.get("src", "")).read_bytes()
            except OSError:
                failed.append({"part": stem, "reason": "file missing after export"})
                dropped.add(dest)
                continue
            if dest.lower().endswith(".pdf") and (quote or e.get("stamp")):
                try:
                    data = _stamp(data, quote)
                except Exception as exc:
                    failed.append({"part": stem, "reason": f"PDF unreadable: {exc}"})
                    dropped.add(dest)
                    continue
                stamped += 1
            try:
                target.parent.mkdir(parents=True, exist_ok=True)
                target.write_bytes(data)
            except OSError as exc:
                failed.append({"part": stem, "reason": f"could not write file: {exc}"})
                dropped.add(dest)
                continue
            written.append(dest)
        rows = []
        for r in m.get("rows") or []:
            r = dict(r)
            r["files"] = [f for f in r.get("files") or [] if f not in dropped]
            rows.append(r)
        counts = {
            "step": sum(d.lower().endswith(".stp") for d in written),
            "pdf": sum(d.startswith("PDF/") for d in written),
            "dxf": sum(d.lower().endswith(".dxf") for d in written),
            "print": sum(d.lower().endswith((".3mf", ".stl")) for d in written),
        }
        all_failed = [dict(f) for f in m.get("failed") or []] + failed
        if partner:
            cover = _cover(m, counts, len(rows), all_failed, quote)
            if quote:
                cover = _stamp(cover, True)
            (staging / "00 Cover.pdf").write_bytes(cover)
            _write_sheet(staging / "BOM.xlsx", m, rows, True, quote)
        else:
            _write_sheet(staging / "Print List.xlsx", m, rows, False, quote)
        total = _zip(staging, out_zip)
        return {"ok": True, "mode": "package", "zip": str(out_zip), "files": total, "rows": len(rows),
                **counts, "stamped": stamped, "failed": len(all_failed),
                "guessed": len(m.get("guessed") or []), "missing": len(m.get("missing_drawings") or []),
                "no_flat_pattern": len(m.get("no_flat_pattern") or []), "left_out": len(m.get("left_out") or [])}
    finally:
        shutil.rmtree(staging, ignore_errors=True)
